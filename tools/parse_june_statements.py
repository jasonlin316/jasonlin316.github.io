from __future__ import annotations

import json
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader


ROOT = Path("/Users/jasonlin316/Documents/2026/June statements")
OUT = Path("/Users/jasonlin316/Documents/GitHub/jasonlin316.github.io/outputs/june_spending_analysis")

PDFS = [
    ROOT / "2026-06-24.pdf",
    ROOT / "20260623-statements-0841-.pdf",
    ROOT / "20260623-statements-4592-.pdf",
    ROOT / "20260623-statements-6910-.pdf",
    ROOT / "eStmt_2026-06-23.pdf",
]


@dataclass
class Txn:
    statement_file: str
    issuer: str
    card: str
    transaction_date: date
    description: str
    amount: float
    transaction_type: str
    category: str

    @property
    def month(self) -> str:
        return self.transaction_date.strftime("%Y-%m")


def pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def money(value: str) -> float:
    value = value.replace("$", "").replace(",", "").strip()
    return float(value)


def as_date(mmdd: str, year: int = 2026) -> date:
    month, day = [int(part) for part in mmdd.split("/")[:2]]
    return date(year, month, day)


def categorize(description: str, transaction_type: str) -> str:
    d = description.upper()
    if transaction_type == "Payment":
        return "Payments"
    if transaction_type == "Credit":
        return "Credits"
    rules = [
        ("Groceries", ["WHOLEFDS", "99 RANCH", "MEGAMART"]),
        ("Restaurants & Coffee", [
            "UBER EATS", "FANTUAN", "RICEBURRITO", "YAYOI", "MARUGAME",
            "REDWOOD BY CHEF", "EIGHTY-EIGHT SUSHI", "MCDONALD", "MARUFUKU", "SUBWAY",
            "CLOCK TOWER", "MOON TEA", "DIN DING", "IKE'S", "SAINT FRANK",
            "FONG KEE", "STARBUCKS", "LINEA CAFFE", "DUMPLING KITCHEN",
            "THE MELT", "KLATCH", "HARU JAPANESE"
        ]),
        ("Pets", ["PETCO", "PETS IN NEED", "PETSMART", "HOMEAGAIN", "UNIVERSAL PET", "CHEWY", "HUMANE SOCIETY"]),
        ("Shopping", ["AMAZON", "TARGET", "HOME DEPOT", "WAYFAIR", "STANFORD SHOPPING", "RALPH LAUREN", "LULULEMON"]),
        ("Insurance", ["LEMONADE", "GEICO", "RENTERS/CONDO"]),
        ("Subscriptions & Software", ["OPENAI", "CHATGPT", "APPLE.COM", "YOUTUBEPREMIUM", "WALMART+", "AMAZON PRIME"]),
        ("Transportation", ["FASTRAK", "PARKING", "CHEVRON"]),
        ("Utilities & Bills", ["CITY OF PALO ALTO"]),
        ("Education & Research", ["PROQUEST"]),
        ("Entertainment", ["WINERY"]),
        ("Misc.", ["BROADWAY TOBACCONISTS", "FANVUE", "ORI FUTURE"]),
    ]
    for category, needles in rules:
        if any(needle in d for needle in needles):
            return category
    return "Misc."


def parse_chase(path: Path, text: str) -> list[Txn]:
    card_match = re.search(r"Account Number:\s+XXXX XXXX XXXX (\d{4})", text)
    card = f"Chase {card_match.group(1) if card_match else path.stem[-5:-1]}"
    lines = text.splitlines()
    txns: list[Txn] = []
    in_table = False
    row_re = re.compile(r"^(\d{2}/\d{2})\s+(.+?)\s+(-?\d[\d,]*\.\d{2}|\.50)$")
    for line in lines:
        if line.startswith("Date of"):
            in_table = True
            continue
        if not in_table:
            continue
        if line.startswith("Total fees charged") or line.startswith("Year-to-date"):
            break
        match = row_re.match(line.strip())
        if not match:
            continue
        tx_date, description, amount_s = match.groups()
        amount = money("0" + amount_s if amount_s.startswith(".") else amount_s)
        transaction_type = "Payment" if amount < 0 or "PAYMENT" in description.upper() else "Purchase"
        category = categorize(description, transaction_type)
        txns.append(Txn(path.name, "Chase", card, as_date(tx_date), description.strip(), amount, transaction_type, category))
    return txns


def parse_amex(path: Path, text: str) -> list[Txn]:
    card_match = re.search(r"Account Ending\s+([\d-]+)", text)
    card = f"Amex {card_match.group(1).split('-')[-1] if card_match else 'unknown'}"
    lines = text.splitlines()
    txns: list[Txn] = []

    def collect(start_idx: int, stop_markers: tuple[str, ...], tx_type: str) -> int:
        i = start_idx
        current_date = None
        desc_parts: list[str] = []
        date_re = re.compile(r"^(\d{2}/\d{2}/\d{2})\*?\s+(.+)$")
        amount_re = re.compile(r"^-?\$[\d,]+\.\d{2}")
        while i < len(lines):
            line = lines[i].strip()
            if any(line.startswith(marker) for marker in stop_markers):
                break
            date_match = date_re.match(line)
            if date_match:
                current_date = date_match.group(1)
                desc_parts = [date_match.group(2).strip()]
            elif current_date and amount_re.match(line):
                amount = money(line.split()[0])
                description = " ".join(desc_parts).strip()
                tx_date = date(2000 + int(current_date[6:8]), int(current_date[0:2]), int(current_date[3:5]))
                category = categorize(description, tx_type)
                txns.append(Txn(path.name, "American Express", card, tx_date, description, amount, tx_type, category))
                current_date = None
                desc_parts = []
            elif current_date and line and not line.startswith(("p. ", "Continued", "Amount")):
                desc_parts.append(line)
            i += 1
        return i

    for idx, line in enumerate(lines):
        if line.strip() == "Payments Amount":
            collect(idx + 1, ("Credits Amount",), "Payment")
        if line.strip() == "Credits Amount":
            collect(idx + 1, ("New Charges",), "Credit")
        if line.strip() == "Amount" and idx > 300:
            collect(idx + 1, ("Fees", "Total Fees"), "Purchase")
    return txns


def parse_bofa(path: Path, text: str) -> list[Txn]:
    card_match = re.search(r"Account#\s+.*?(\d{4})\s+May", text)
    card = f"Bank of America {card_match.group(1) if card_match else 'unknown'}"
    match = re.search(
        r"Payments and Other Credits(?P<payments>.+?)TOTAL PAYMENTS AND OTHER CREDITS FOR THIS PERIOD .+?"
        r"Purchases and Adjustments(?P<purchases>.+?)TOTAL PURCHASES AND ADJUSTMENTS FOR THIS PERIOD",
        text,
        flags=re.S,
    )
    if not match:
        return []
    txns: list[Txn] = []
    row_re = re.compile(r"(\d{2}/\d{2})\s+(\d{2}/\d{2})\s+(.+?)\s+\d{4}\s+\d{4}\s+(-?\d[\d,]*\.\d{2})")
    for block_name, tx_type in [("payments", "Payment"), ("purchases", "Purchase")]:
        block = match.group(block_name)
        for row in row_re.finditer(block):
            tx_date, _posting_date, description, amount_s = row.groups()
            amount = money(amount_s)
            transaction_type = tx_type if amount >= 0 else "Payment"
            category = categorize(description, transaction_type)
            txns.append(Txn(path.name, "Bank of America", card, as_date(tx_date), description.strip(), amount, transaction_type, category))
    return txns


def parse_all() -> list[Txn]:
    txns: list[Txn] = []
    for path in PDFS:
        text = pdf_text(path)
        if path.name.startswith("20260623-statements"):
            txns.extend(parse_chase(path, text))
        elif path.name == "2026-06-24.pdf":
            txns.extend(parse_amex(path, text))
        elif path.name.startswith("eStmt"):
            txns.extend(parse_bofa(path, text))
    unique: list[Txn] = []
    seen = set()
    for txn in txns:
        key = (
            txn.statement_file,
            txn.transaction_date.isoformat(),
            txn.description,
            round(txn.amount, 2),
            txn.transaction_type,
        )
        if txn.issuer == "American Express" and key in seen:
            continue
        seen.add(key)
        unique.append(txn)
    return unique


def write_outputs(txns: list[Txn]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    df = pd.DataFrame([t.__dict__ | {"month": t.month} for t in txns])
    df = df[[
        "statement_file", "issuer", "card", "transaction_date", "month",
        "description", "amount", "transaction_type", "category",
    ]]
    df = df.sort_values(["transaction_date", "issuer", "card", "description"]).reset_index(drop=True)
    csv_path = OUT / "june_statement_transactions.csv"
    xlsx_path = OUT / "june_spending_analysis.xlsx"
    df.to_csv(csv_path, index=False)

    purchases = df[df["transaction_type"] == "Purchase"].copy()
    by_category = purchases.groupby("category", as_index=False)["amount"].sum().sort_values("amount", ascending=False)
    by_month = purchases.groupby("month", as_index=False)["amount"].sum()
    by_card = purchases.groupby(["issuer", "card"], as_index=False)["amount"].sum().sort_values("amount", ascending=False)
    by_month_category = purchases.pivot_table(index="month", columns="category", values="amount", aggfunc="sum", fill_value=0).reset_index()

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        overview_rows = [
            ["Metric", "Value"],
            ["Total gross purchases in these June statements", purchases["amount"].sum()],
            ["Total credits/refunds", df[df["transaction_type"] == "Credit"]["amount"].sum()],
            ["Payments excluded from spending totals", df[df["transaction_type"] == "Payment"]["amount"].sum()],
            ["Transaction count", len(df)],
            ["Purchase transaction count", len(purchases)],
            ["Statements parsed", len(PDFS)],
        ]
        pd.DataFrame(overview_rows[1:], columns=overview_rows[0]).to_excel(writer, sheet_name="Overview", index=False)
        by_category.to_excel(writer, sheet_name="By Category", index=False)
        by_month.to_excel(writer, sheet_name="By Month", index=False)
        by_card.to_excel(writer, sheet_name="By Card", index=False)
        by_month_category.to_excel(writer, sheet_name="Month x Category", index=False)
        df.to_excel(writer, sheet_name="Transactions", index=False)

    wb = load_workbook(xlsx_path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F5597")
            cell.alignment = Alignment(horizontal="center")
        for col_idx, column in enumerate(ws.columns, start=1):
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 48)
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, (int, float)) and ("amount" in str(ws.cell(1, cell.column).value).lower() or cell.column == 2):
                    cell.number_format = '$#,##0.00;[Red]-$#,##0.00'

    ws = wb["By Category"]
    chart = BarChart()
    chart.title = "Gross Purchases by Category"
    chart.y_axis.title = "Amount"
    chart.x_axis.title = "Category"
    data = Reference(ws, min_col=2, min_row=1, max_row=ws.max_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 8
    chart.width = 14
    ws.add_chart(chart, "D2")

    wb.save(xlsx_path)

    summary = {
        "transactions": len(df),
        "purchases": len(purchases),
        "gross_purchases": round(float(purchases["amount"].sum()), 2),
        "credits": round(float(df[df["transaction_type"] == "Credit"]["amount"].sum()), 2),
        "payments": round(float(df[df["transaction_type"] == "Payment"]["amount"].sum()), 2),
        "by_category": {r["category"]: round(float(r["amount"]), 2) for _, r in by_category.iterrows()},
        "by_month": {r["month"]: round(float(r["amount"]), 2) for _, r in by_month.iterrows()},
        "by_card": {f'{r["card"]}': round(float(r["amount"]), 2) for _, r in by_card.iterrows()},
    }
    (OUT / "summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")


if __name__ == "__main__":
    transactions = parse_all()
    write_outputs(transactions)
    print(f"Parsed {len(transactions)} transactions")
